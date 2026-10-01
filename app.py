import io
import json
import os
import re
import sqlite3
from datetime import date
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from groq import Groq
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

load_dotenv()
HERE = Path(__file__).parent
DB = HERE / "shop.db"
MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
groq = Groq(api_key=os.getenv("GROQ_API_KEY"))
app = FastAPI(title="Fruit Shop Ledger")


def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con


with db() as con:
    con.execute(
        """CREATE TABLE IF NOT EXISTS entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            day TEXT, kind TEXT, fruit TEXT,
            qty REAL, unit TEXT, price REAL, total REAL, note TEXT, place TEXT
        )"""
    )
    # keep old data: add the "place" column if the table came from the first version
    cols = [r["name"] for r in con.execute("PRAGMA table_info(entries)")]
    if "place" not in cols:
        con.execute("ALTER TABLE entries ADD COLUMN place TEXT")

PARSE_PROMPT = (
    "You extract one record for a fruit shop from the user's sentence. The "
    "sentence may be in English, Hindi or Telugu. Reply with ONLY a JSON object, "
    'no other text: {"kind": "sale" or "purchase", "fruit": "English name", '
    '"qty": number, "unit": "kg" or "dozen" or "piece", '
    '"price": number (price per unit), '
    '"place": string (customer, hotel or supplier name, e.g. "Hotel 1"), '
    '"note": string}. '
    "Use null for anything not mentioned. If the user does not say bought or "
    'purchased, kind is "sale".'
)


class Text(BaseModel):
    text: str


class Entry(BaseModel):
    kind: str = "sale"
    fruit: str
    qty: float
    unit: str = "kg"
    price: float
    place: str = ""
    note: str = ""


@app.get("/")
def home():
    return FileResponse(HERE / "index.html")


@app.post("/parse")
def parse(body: Text):
    if not body.text.strip():
        raise HTTPException(400, "Nothing was heard.")
    try:
        resp = groq.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": PARSE_PROMPT},
                {"role": "user", "content": body.text},
            ],
        )
        raw = resp.choices[0].message.content or ""
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        return json.loads(match.group(0))
    except Exception as e:
        raise HTTPException(502, f"Could not understand that: {e}")


@app.post("/entries")
def add_entry(e: Entry):
    kind = "purchase" if e.kind == "purchase" else "sale"
    total = round(e.qty * e.price, 2)
    with db() as con:
        con.execute(
            "INSERT INTO entries (day, kind, fruit, qty, unit, price, total, note, place) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (date.today().isoformat(), kind, e.fruit.strip().title(), e.qty,
             e.unit, e.price, total, e.note, e.place.strip().title()),
        )
    return {"ok": True}


@app.get("/entries")
def list_entries():
    with db() as con:
        rows = [dict(r) for r in con.execute(
            "SELECT * FROM entries ORDER BY id DESC LIMIT 50")]
        places = [r["place"] for r in con.execute(
            "SELECT DISTINCT place FROM entries WHERE place != '' AND place IS NOT NULL")]
    return {"entries": rows, "places": places}


@app.delete("/entries/{entry_id}")
def delete_entry(entry_id: int):
    with db() as con:
        con.execute("DELETE FROM entries WHERE id=?", (entry_id,))
    return {"ok": True}


def build_report(period: str):
    """Profit = sales - (average buying price x quantity sold)."""
    today = date.today().isoformat()
    with db() as con:
        rows = [dict(r) for r in con.execute("SELECT * FROM entries ORDER BY id")]

    def in_period(d):
        if period == "today":
            return d == today
        if period == "month":
            return d[:7] == today[:7]
        return True

    # average buying price per fruit+unit, from ALL purchases
    bought = {}
    for r in rows:
        if r["kind"] == "purchase":
            b = bought.setdefault((r["fruit"], r["unit"]), [0.0, 0.0])
            b[0] += r["qty"]
            b[1] += r["total"]
    avg = {k: v[1] / v[0] for k, v in bought.items() if v[0]}

    fruits, places = {}, {}
    spent = 0.0
    for r in rows:
        key = (r["fruit"], r["unit"])
        f = fruits.setdefault(key, dict(
            fruit=r["fruit"], unit=r["unit"], stock=0.0, sold=0.0,
            revenue=0.0, cost=0.0, cost_known=True))
        is_buy = r["kind"] == "purchase"
        f["stock"] += r["qty"] if is_buy else -r["qty"]   # stock is all-time
        if not in_period(r["day"]):
            continue
        if is_buy:
            spent += r["total"]
            continue
        f["sold"] += r["qty"]
        f["revenue"] += r["total"]
        if key in avg:
            cost = avg[key] * r["qty"]
        else:
            cost = 0.0
            f["cost_known"] = False
        f["cost"] += cost
        p = places.setdefault(r["place"] or "(no name)", dict(revenue=0.0, known=0.0, cost=0.0))
        p["revenue"] += r["total"]
        if key in avg:
            p["known"] += r["total"]
        p["cost"] += cost

    by_fruit = []
    for f in fruits.values():
        if f["sold"] or f["stock"]:
            f["profit"] = round(f["revenue"] - f["cost"], 2) if f["cost_known"] else None
            by_fruit.append({k: round(v, 2) if isinstance(v, float) else v for k, v in f.items()})
    by_fruit.sort(key=lambda x: -x["revenue"])
    by_place = sorted(
        [dict(place=k, revenue=round(v["revenue"], 2),
              profit=round(v["known"] - v["cost"], 2)) for k, v in places.items()],
        key=lambda x: -x["revenue"])
    revenue = sum(f["revenue"] for f in fruits.values())
    known_rev = sum(f["revenue"] for f in fruits.values() if f["cost_known"])
    cost = sum(f["cost"] for f in fruits.values())
    return {
        "sales": round(revenue, 2),
        "purchases": round(spent, 2),
        "profit": round(known_rev - cost, 2),
        "by_fruit": by_fruit,
        "by_place": by_place,
    }


@app.get("/report")
def report(period: str = "today"):
    return build_report(period)


@app.get("/export")
def export():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["Date", "Type", "Fruit", "Customer / place", "Qty", "Unit",
               "Price per unit", "Total", "Note"])
    with db() as con:
        for r in con.execute("SELECT * FROM entries ORDER BY id"):
            ws.append([r["day"], r["kind"], r["fruit"], r["place"], r["qty"],
                       r["unit"], r["price"], r["total"], r["note"]])

    rep = build_report("all")
    w2 = wb.create_sheet("Profit by fruit")
    w2.append(["Fruit", "Unit", "Sold qty", "Sales", "Cost of sold", "Profit / Loss", "Stock left"])
    for f in rep["by_fruit"]:
        w2.append([f["fruit"], f["unit"], f["sold"], f["revenue"], f["cost"], f["profit"], f["stock"]])
    w3 = wb.create_sheet("By customer")
    w3.append(["Customer / place", "Sales", "Profit / Loss"])
    for p in rep["by_place"]:
        w3.append([p["place"], p["revenue"], p["profit"]])
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="fruit_shop_ledger.xlsx"'},
    )